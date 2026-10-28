"""Computational verification against independent formulas for invented instruments.

This is not human review, psychometric validation or a clinical effectiveness study.
"""
import csv,io,itertools,json,math,random,tempfile,unittest
from fractions import Fraction
from copy import deepcopy
from pathlib import Path
from qreport import ReportEngine,ConfigurationError,DataError,render_html,load_config,score_rows,main
from make_examples import configurations,bands

COUNTS={'exhaustive_response_profiles':0,'scored_scale_comparisons':0,'unscorable_profiles_checked':0,
        'count_rule_profiles':0,'lookup_records':0,'computed_item_records':0,'dynamic_text_records':0}

def record(values):return dict(participant_id='TEST001',**{f'q{i+1}':v for i,v in enumerate(values)})

class ScoringChecks(unittest.TestCase):
    def test_exhaustive_sum_mean_and_reverse(self):
        e=ReportEngine(configurations()['likert'])
        for values in itertools.product(range(5),repeat=4):
            r=e.score(record(values));expected=values[0]+(4-values[1])+values[2]+values[3]
            self.assertEqual(r['scales'][0]['score'],expected)
            self.assertEqual(r['scales'][1]['score'],expected/4)
            self.assertEqual(r['scales'][0]['band']['id'],'lower' if expected<8 else 'upper')
            self.assertEqual(len(r['sections']),int(expected>=8 and values[0]<=1))
            COUNTS['exhaustive_response_profiles']+=1;COUNTS['scored_scale_comparisons']+=2

    def test_exhaustive_weighted_mean_and_recoding(self):
        e=ReportEngine(configurations()['weighted_mean']);labels=['never','sometimes','often','always']
        for v in itertools.product(range(4),repeat=4):
            r=e.score(record([labels[x] for x in v]));expected=(v[0]+2*v[1]+v[2]+2*v[3])/6
            self.assertAlmostEqual(r['scales'][0]['score'],expected,places=12)
            self.assertEqual(r['scales'][0]['band']['id'],'lower' if expected<1.5 else 'upper')
            COUNTS['exhaustive_response_profiles']+=1;COUNTS['scored_scale_comparisons']+=1

    def test_exhaustive_binary_conversion(self):
        e=ReportEngine(configurations()['knowledge'])
        for v in itertools.product(['A','B'],repeat=5):
            r=e.score(record(v));expected=20*v.count('A')
            self.assertEqual(r['scales'][0]['score'],expected)
            self.assertEqual(r['scales'][0]['band']['id'],'lower' if expected<60 else 'upper')
            COUNTS['exhaustive_response_profiles']+=1;COUNTS['scored_scale_comparisons']+=1

    def test_exhaustive_proration_and_missingness(self):
        e=ReportEngine(configurations()['prorated']);w=[1,2,1,2]
        for v in itertools.product([None,1,2,3,4,5],repeat=4):
            r=e.score(record(v));s=r['scales'][0];n=sum(x is not None for x in v)
            if n<3:
                self.assertEqual(r['status'],'unscored');self.assertIsNone(s['score']);COUNTS['unscorable_profiles_checked']+=1
            else:
                x=[(6-z if j==1 else z) if z is not None else None for j,z in enumerate(v)]
                expected=sum(wj*z for wj,z in zip(w,x) if z is not None)*6/sum(wj for wj,z in zip(w,x) if z is not None)
                self.assertAlmostEqual(s['score'],expected,places=12)
                self.assertEqual(s['prorated'],n<4)
                COUNTS['scored_scale_comparisons']+=1
            COUNTS['exhaustive_response_profiles']+=1

    def test_zero_is_a_valid_response(self):
        r=ReportEngine(configurations()['likert']).score(record([0,0,0,0]))
        self.assertEqual(r['scales'][0]['score'],4);self.assertEqual(r['scales'][0]['answered'],4)

    def test_continuous_response_and_affine_conversion(self):
        c=configurations()['likert'];c['items'][0]['allow_decimal']=True
        c['scales']=c['scales'][:1];c['scales'][0].update(multiplier=1.5,offset=2,bands=bands(2,14,26))
        c['sections']=[]
        r=ReportEngine(c).score(record([.5,2,3,4]))
        self.assertEqual(r['scales'][0]['score'],16.25)
        self.assertEqual(r['scales'][0]['band']['id'],'upper')

    def test_word_export_contains_the_scored_values(self):
        from qreport import render_docx
        from docx import Document
        r=ReportEngine(configurations()['likert']).score(record([1,2,3,4]))
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'report.docx';render_docx(r,path);d=Document(path)
            self.assertEqual(d.tables[0].cell(1,1).text,'10')
            self.assertEqual(d.tables[0].cell(2,1).text,'2.5')
            self.assertIn(r['config_sha256'],'\n'.join(p.text for p in d.paragraphs))

    def test_invalid_values_block_release(self):
        e=ReportEngine(configurations()['likert'])
        for bad in ['bogus','NaN','Infinity',999,-1,0.5]:
            r=e.score(record([bad,0,0,0]));self.assertEqual(r['status'],'blocked')
            with self.assertRaises(DataError):render_html(r)

    def test_all_declared_missing_codes(self):
        e=ReportEngine(configurations()['likert'])
        for missing in ['',None,'NA',99,'99.0']:
            self.assertEqual(e.score(record([missing,0,0,0]))['status'],'unscored')

    def test_column_order_and_extra_metadata(self):
        e=ReportEngine(configurations()['likert']);r=record([1,2,3,4])
        reordered=dict(reversed(list(r.items())));reordered['unused']='ignored'
        self.assertEqual(e.score(r),e.score(reordered))

    def test_absent_columns_are_not_missing_answers(self):
        e=ReportEngine(configurations()['likert']);r=record([1,2,3,4]);del r['q2']
        with self.assertRaises(DataError):e.score(r)

    def test_subscale_partial_report(self):
        c=configurations()['likert'];c['scales'][1]['items']=['q1','q3']
        r=ReportEngine(c).score(record([1,None,3,4]))
        self.assertEqual(r['status'],'partial');self.assertIn('Not scored',render_html(r))

    def test_unknown_condition_is_not_true_under_not(self):
        c=configurations()['likert'];c['sections'][0]['when']={'not':{'item':'q2','op':'gt','value':1}}
        r=ReportEngine(c).score(record([1,None,3,4]))
        self.assertEqual(r['section_trace'][0]['decision'],'withheld_unknown');self.assertEqual(r['sections'],[])

    def test_short_csv_row_stops_the_batch(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);cf=td/'config.json';cf.write_text(json.dumps(configurations()['prorated']))
            df=td/'input.csv';df.write_text('participant_id,q1,q2,q3,q4\nP1,1,2,3,4\nP2,1,2,3\n')
            with self.assertRaises(DataError):main(['--config',str(cf),'--data',str(df),'--out',str(td/'reports')])
            self.assertFalse((td/'reports').exists())

    def test_displayed_scores_are_rounded_and_decisions_use_exact_scores(self):
        c=configurations()['weighted_mean'];e=ReportEngine(c)
        r=e.score(record(['sometimes','never','never','never']));s=r['scales'][0]
        self.assertEqual(s['score_exact'],'0.1666666666666666666666666667');self.assertEqual(s['score_display'],'0.17')
        self.assertIn('score is 0.17.',s['text']);self.assertNotIn('0.1666',render_html(r))
        c['scales'][0]['display_decimals']=0;s=ReportEngine(c).score(record(['always','always','always','often']))['scales'][0]
        self.assertEqual(s['score_display'],'3');self.assertEqual(s['band']['id'],'upper')
        self.assertEqual(ReportEngine(configurations()['likert']).score(record([1,2,3,4]))['scales'][0]['score_display'],'10')
        for bad in [True,-1,7,1.5]:
            c=configurations()['weighted_mean'];c['scales'][0]['display_decimals']=bad
            with self.assertRaises(ConfigurationError):ReportEngine(c)

    def test_condition_operators_and_three_valued_logic(self):
        c=configurations()['likert']
        ops={'eq':lambda a,b:a==b,'ne':lambda a,b:a!=b,'lt':lambda a,b:a<b,'le':lambda a,b:a<=b,'gt':lambda a,b:a>b,'ge':lambda a,b:a>=b}
        c['sections']=[dict(id=f'{op}_{v}',title='t',text='x',source='s',when={'item':'q1','op':op,'value':v}) for op in ops for v in range(5)]
        e=ReportEngine(c)
        for q1 in range(5):
            got={t['id']:t['decision'] for t in e.score(record([q1,0,0,0]))['section_trace']}
            for op,f in ops.items():
                for v in range(5):self.assertEqual(got[f'{op}_{v}'],'included' if f(q1,v) else 'excluded')
        # Kleene truth tables: leaves that are true, false or unknown (q2 missing), combined under all, any and not
        leaf={True:{'item':'q1','op':'ge','value':0},False:{'item':'q1','op':'lt','value':0},None:{'item':'q2','op':'ge','value':0}}
        k_and=lambda a,b:False if False in (a,b) else None if None in (a,b) else True
        k_or=lambda a,b:True if True in (a,b) else None if None in (a,b) else False
        k_not=lambda a:None if a is None else not a
        name={True:'T',False:'F',None:'U'};expected={};c['sections']=[]
        for a in (True,False,None):
            for b in (True,False,None):
                for kind,node,want in (('all',{'all':[leaf[a],leaf[b]]},k_and(a,b)),('any',{'any':[leaf[a],leaf[b]]},k_or(a,b)),
                                       ('notall',{'not':{'all':[leaf[a],leaf[b]]}},k_not(k_and(a,b)))):
                    sid=f'{kind}_{name[a]}{name[b]}';expected[sid]=want
                    c['sections'].append(dict(id=sid,title='t',text='x',source='s',when=node))
        got={t['id']:t['decision'] for t in ReportEngine(c).score(record([2,None,0,0]))['section_trace']}
        for sid,want in expected.items():
            self.assertEqual(got[sid],{True:'included',False:'excluded',None:'withheld_unknown'}[want],sid)

    def test_proration_of_a_weighted_mean(self):
        c=configurations()['prorated'];c['scales'][0].update(method='weighted_mean',bands=bands(1,3,5))
        s=ReportEngine(c).score(record([5,None,1,3]))['scales'][0]
        self.assertEqual(s['score_exact'],'3');self.assertTrue(s['prorated']);self.assertEqual(s['band']['id'],'upper')
        s=ReportEngine(c).score(record([5,4,1,3]))['scales'][0]   # q2 is reverse-scored: 6 - 4 = 2
        self.assertAlmostEqual(s['score'],(5+2*2+1+2*3)/6,places=12);self.assertFalse(s['prorated']);self.assertEqual(s['band']['id'],'lower')

    def test_eligibility_and_manual_review_hold(self):
        c=configurations()['likert'];c['eligibility']=[dict(column='complete',equals=2,reason='Incomplete form')]
        e=ReportEngine(c);r=record([1,2,3,4]);r['complete']=1
        self.assertEqual(e.score(r)['status'],'blocked')
        c['scales'][0]['bands'][1]['hold']=True;e=ReportEngine(c);r['complete']=2
        result=e.score(r);self.assertEqual(result['status'],'held_for_review')
        with self.assertRaises(DataError):render_html(result)

    def test_duplicate_missing_and_unsafe_ids(self):
        e=ReportEngine(configurations()['likert']);r=record([1,2,3,4])
        with self.assertRaises(DataError):score_rows(e,[r,r])
        for value in [None,'','a\nb']:
            with self.assertRaises(DataError):e.score(dict(r,participant_id=value))

    def test_config_errors(self):
        changes=[lambda c:c.update(unknown_field=True),
                 lambda c:c['items'][0].update(reverse='yes'),
                 lambda c:c['scales'][0]['bands'][1].update(lower=9),
                 lambda c:c['scales'][0]['bands'][1].update(lower=7),
                 lambda c:c['scales'][0].update(items=['unknown']),
                 lambda c:c['scales'][0].update(weights={'q1':2}),
                 lambda c:c.update(missing_codes=[0]),
                 lambda c:c['items'][0].update(column='q2'),
                 lambda c:c['scales'][0]['bands'][0].update(text='{{unknown}}'),
                 lambda c:c['sections'][0].update(when={'scale':'missing','op':'gt','value':1})]
        for change in changes:
            c=configurations()['likert'];change(c)
            with self.assertRaises(ConfigurationError):ReportEngine(c)

    def test_proration_requires_explicit_rules(self):
        c=configurations()['prorated'];del c['scales'][0]['missing_policy']['justification']
        with self.assertRaises(ConfigurationError):ReportEngine(c)
        c=configurations()['prorated'];c['items'][0]['max']=6
        with self.assertRaises(ConfigurationError):ReportEngine(c)

    def test_duplicate_json_keys_are_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'bad.json';p.write_text('{"schema_version":"1.0","schema_version":"2.0"}')
            with self.assertRaises(ConfigurationError):load_config(p)

    def test_decisions_are_deterministic_and_hash_tracks_rules(self):
        c=configurations()['likert'];e=ReportEngine(c);r=record([1,2,3,4])
        self.assertEqual(e.score(r),e.score(r))
        self.assertEqual(e.config_sha256,ReportEngine(json.loads(json.dumps(c))).config_sha256)
        c['instrument_version']='changed';self.assertNotEqual(e.config_sha256,ReportEngine(c).config_sha256)

    def test_html_escapes_text_and_preserves_provenance(self):
        c=configurations()['likert'];c['title']='<script>bad</script>'
        r=ReportEngine(c).score(record([1,2,3,4]));h=render_html(r)
        self.assertNotIn('<script>',h);self.assertIn('&lt;script&gt;',h)
        self.assertIn(r['config_sha256'],h);self.assertIn('10',h);self.assertIn('Example conditional message',h)

    def test_csv_batch_outputs_and_overwrite_guard(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);cf=td/'config.json';cf.write_text(json.dumps(configurations()['likert']))
            df=td/'input.csv';df.write_text('participant_id,q1,q2,q3,q4\n../TEST,1,2,3,4\nBLOCKED,1,2,999,4\n')
            out=td/'reports'
            import contextlib
            with contextlib.redirect_stdout(io.StringIO()):main(['--config',str(cf),'--data',str(df),'--out',str(out)])
            m=json.loads((out/'run_manifest.json').read_text())
            self.assertEqual(m['status_counts'],{'ready':1,'blocked':1})
            self.assertEqual(len(list(out.glob('*.html'))),1)
            self.assertTrue(all('/' not in p.stem for p in out.iterdir()))
            with self.assertRaises(DataError):main(['--config',str(cf),'--data',str(df),'--out',str(out)])



def row_of(header, values):
    return {h: ('' if v is None else str(v)) for h, v in zip(header, values)}


def half_up(x):
    """Round a nonnegative Fraction half up (written separately from the engine's Decimal rounding)."""
    return math.floor(x + Fraction(1, 2))


class ExtensionChecks(unittest.TestCase):
    """Schema 1.1 features, each compared with a calculation written independently in this file."""

    # ---------------------------------------------------------------- count rules
    def test_exhaustive_count_rule_and_combined_condition(self):
        e = ReportEngine(configurations()['count_rule'])
        for v in itertools.product(range(4), repeat=6):
            r = e.score(record(v))
            count = sum(x >= 2 for x in v[:5]) + (v[5] >= 1)
            self.assertEqual(r['scales'][0]['score_exact'], str(count))
            self.assertEqual(r['scales'][0]['band']['id'], 'lower' if count < 4 else 'upper')
            self.assertEqual(r['scales'][1]['score_exact'], str(sum(v)))
            pattern = count >= 4 and (v[0] >= 2 or v[1] >= 2)
            self.assertEqual(r['section_trace'][0]['decision'], 'included' if pattern else 'excluded')
            COUNTS['count_rule_profiles'] += 1
        r = e.score(record([3, None, 3, 3, 3, 3]))   # count needs every item; item 1 alone already satisfies the 'any'
        self.assertEqual(r['scales'][0]['status'], 'unscored'); self.assertEqual(r['section_trace'][0]['decision'], 'withheld_unknown')

    def test_count_configuration_errors(self):
        changes = [lambda c: c['scales'][0].pop('criterion'),
                   lambda c: c['scales'][0].update(missing_policy={'method': 'prorate', 'minimum_answered': 5, 'justification': 'x'}),
                   lambda c: c['scales'][0].update(item_criteria={'zz': {'op': 'ge', 'value': 1}}),
                   lambda c: c['scales'][0].update(criterion={'op': 'between', 'value': 1}),
                   lambda c: c['scales'][1].update(criterion={'op': 'ge', 'value': 1})]
        for change in changes:
            c = configurations()['count_rule']; change(c)
            with self.assertRaises(ConfigurationError): ReportEngine(c)

    # ---------------------------------------------------------------- lookup tables
    def test_exhaustive_norm_table_lookup(self):
        from make_examples import NORM_GROUPS
        e = ReportEngine(configurations()['norm_table'])
        header = ['participant_id', 'age', 'sex', 'q1', 'q2', 'q3', 'q4']
        groups = [(34, 'F'), (34, 'M'), (16, 'F'), (16, 'M'), (16, None), (40, None), (None, 'F'), (17, 'M'), (18, 'M')]
        for (age, sex), v in itertools.product(groups, itertools.product([None, 1, 2, 3, 4, 5], repeat=4)):
            r = e.score(row_of(header, ['T1', age, sex, *v])); s = r['scales'][0]
            answered = [x for x in v if x is not None]
            if age is None or (age >= 18 and sex is None):
                gid = None
            else:
                gid = 'youth' if age < 18 else 'adult_female' if sex == 'F' else 'adult_male'
            if len(answered) < 3:
                self.assertEqual(s['status'], 'unscored'); self.assertIn('Too few', s['reason'])
            elif gid is None:
                self.assertEqual(s['status'], 'unscored'); self.assertIn('grouping variable', s['reason'])
            else:
                raw = math.ceil(Fraction(sum(answered) * 4, len(answered)))
                mean, sd = NORM_GROUPS[gid]
                t = half_up(50 + 10 * (raw - Fraction(str(mean))) / Fraction(str(sd)))
                self.assertEqual((s['lookup']['table'], s['lookup']['raw_used'], s['score_exact']), (gid, str(raw), str(t)))
                self.assertEqual(s['band']['id'], 'lower' if t < 60 else 'upper')
                self.assertEqual(s['prorated'], len(answered) < 4)
            COUNTS['lookup_records'] += 1

    def interval_config(self, **lookup):
        c = configurations()['norm_table']; s = c['scales'][0]
        s.update(method='mean', missing_policy={'method': 'complete'}, display_decimals=2,
                 lookup=dict(reference='Invented interval table.', **lookup), bands=[])
        s.pop('interpretation_reference')
        return c

    def test_interval_table_boundaries_and_half_up_rounding(self):
        rows = [dict(lower=1, upper=2, score=10), dict(lower=2, upper=3, score=20), dict(lower=3, upper=4, score=30),
                dict(lower=4, upper=5, upper_inclusive=True, score=40)]
        e = ReportEngine(self.interval_config(tables=[dict(id='all', type='interval', rows=rows)]))
        header = ['participant_id', 'age', 'sex', 'q1', 'q2', 'q3', 'q4']
        for v in itertools.product(range(1, 6), repeat=4):
            mean = Fraction(sum(v), 4)
            expected = 40 if mean == 5 else 10 * math.floor(mean)
            self.assertEqual(e.score(row_of(header, ['T', 30, 'F', *v]))['scales'][0]['score_exact'], str(expected))
            COUNTS['lookup_records'] += 1
        exact_rows = [dict(raw=r, score=100 + r) for r in range(1, 6)]
        e = ReportEngine(self.interval_config(round_raw='half_up', tables=[dict(id='all', type='exact', rows=exact_rows)]))
        for v in itertools.product(range(1, 6), repeat=4):
            used = half_up(Fraction(sum(v), 4))
            s = e.score(row_of(header, ['T', 30, 'F', *v]))['scales'][0]
            self.assertEqual((s['lookup']['raw_used'], s['score_exact']), (str(used), str(100 + used)))
            COUNTS['lookup_records'] += 1

    def test_lookup_no_matching_table(self):
        c = configurations()['norm_table']; c['scales'][0]['lookup']['tables'] = c['scales'][0]['lookup']['tables'][:2]
        header = ['participant_id', 'age', 'sex', 'q1', 'q2', 'q3', 'q4']
        s = ReportEngine(c).score(row_of(header, ['T', 15, 'F', 3, 3, 3, 3]))['scales'][0]
        self.assertEqual(s['status'], 'unscored'); self.assertIn('No conversion table', s['reason'])

    def test_lookup_configuration_errors(self):
        def table(c): return c['scales'][0]['lookup']['tables'][0]
        changes = [lambda c: table(c)['rows'].pop(5),                                               # gap
                   lambda c: table(c)['rows'].append(dict(table(c)['rows'][0])),                     # duplicate raw
                   lambda c: table(c).update(type='lookup'),
                   lambda c: c['scales'][0]['lookup'].update(round_raw='none'),                     # prorated raw is not whole
                   lambda c: c['scales'][0].update(multiplier=2),
                   lambda c: table(c).update(group={'item': 'q1', 'op': 'ge', 'value': 1}),         # groups use variables only
                   lambda c: table(c).update(group={'variable': 'sex', 'op': 'eq', 'value': 'X'}),  # unknown category
                   lambda c: table(c).update(group={'variable': 'sex', 'op': 'gt', 'value': 'F'}),
                   lambda c: c['variables'][0].update(max=120),                                      # 99 is a missing code
                   lambda c: c['variables'].append(dict(id='age2', column='q1', type='numeric', min=0, max=5)),
                   lambda c: c.update(schema_version='1.0')]
        for change in changes:
            c = configurations()['norm_table']; change(c)
            with self.assertRaises(ConfigurationError): ReportEngine(c)
        c = self.interval_config(tables=[dict(id='a', type='interval', rows=[dict(lower=1, upper=3, score=1), dict(lower=3, upper=4.5, upper_inclusive=True, score=2)])])
        with self.assertRaises(ConfigurationError): ReportEngine(c)                                  # does not reach 5

    # ---------------------------------------------------------------- computed items
    def test_sleep_diary_against_independent_calculation(self):
        e = ReportEngine(configurations()['sleep_diary']); rng = random.Random(20261028)
        header = ['participant_id', 'bedtime', 'waketime', 'latency_min', 'sleep_hours', 'q1', 'q2']
        maybe = lambda x: None if rng.random() < .08 else x
        for n in range(3000):
            bed, wake = rng.randrange(1440), rng.randrange(1440)
            if n % 100 == 0: wake = bed                                                  # zero minutes in bed
            hours = maybe(Fraction(rng.randrange(0, 65), 4)); lat = maybe(rng.randrange(0, 241))
            q = [maybe(rng.randrange(4)), maybe(rng.randrange(4))]
            values = ['S', f'{bed // 60}:{bed % 60:02d}', f'{wake // 60:02d}:{wake % 60:02d}', lat,
                      None if hours is None else str(float(hours)), *q]
            r = e.score(row_of(header, values)); s = r['scales'][0]
            in_bed = (wake - bed) % 1440
            if hours is not None and in_bed == 0:
                self.assertEqual(r['status'], 'blocked'); self.assertIn('Division by zero', r['reasons'][0])
                COUNTS['computed_item_records'] += 1; continue
            comps = []
            comps.append(None if hours is None else 3 if hours < 5 else 2 if hours < 6 else 1 if hours <= 7 else 0)
            comps.append(None if lat is None else 0 if lat <= 15 else 1 if lat <= 30 else 2 if lat <= 60 else 3)
            if hours is None:
                comps.append(None)
            else:
                eff = hours * 60 / in_bed * 100
                comps.append(3 if eff < 65 else 2 if eff < 75 else 1 if eff < 85 else 0)
            comps += q
            self.assertEqual(r['item_trace']['in_bed']['scored'], str(in_bed))
            if None in comps:
                self.assertEqual(s['status'], 'unscored')
            else:
                self.assertEqual(s['score_exact'], str(sum(comps)))
                self.assertEqual(s['band']['id'], 'lower' if sum(comps) < 6 else 'upper')
            COUNTS['computed_item_records'] += 1

    def test_bin_boundaries_and_clock_wraparound(self):
        e = ReportEngine(configurations()['sleep_diary'])
        base = dict(participant_id='B', bedtime='23:00', waketime='07:00', latency_min='15', q1='0', q2='0')
        for hours, comp in [('4.99', '3'), ('5', '2'), ('5.99', '2'), ('6', '1'), ('7', '1'), ('7.01', '0')]:
            self.assertEqual(e.score(dict(base, sleep_hours=hours))['item_trace']['c_duration']['scored'], comp)
        for lat, comp in [('15', '0'), ('16', '1'), ('30', '1'), ('31', '2'), ('60', '2'), ('61', '3')]:
            self.assertEqual(e.score(dict(base, latency_min=lat, sleep_hours='7'))['item_trace']['c_latency']['scored'], comp)
        for bed, wake, minutes in [('23:00', '07:00', '480'), ('00:00', '23:59', '1439'), ('23:59', '00:00', '1'), ('7:05', '07:04', '1439')]:
            self.assertEqual(e.score(dict(base, bedtime=bed, waketime=wake, sleep_hours='7'))['item_trace']['in_bed']['scored'], minutes)

    def test_invalid_inputs_to_computed_items_block_release(self):
        e = ReportEngine(configurations()['sleep_diary'])
        base = dict(participant_id='B', bedtime='23:00', waketime='07:00', latency_min='15', sleep_hours='7', q1='0', q2='0')
        for change in [dict(bedtime='24:00'), dict(waketime='7.30'), dict(bedtime='noon'), dict(latency_min='241'), dict(sleep_hours='-1')]:
            r = e.score(dict(base, **change)); self.assertEqual(r['status'], 'blocked', change)
            with self.assertRaises(DataError): render_html(r)
        r = e.score(dict(base, bedtime=''))                                              # missing, not invalid
        self.assertEqual(r['status'], 'unscored'); self.assertEqual(r['item_trace']['c_efficiency']['status'], 'missing')
        c = configurations()['sleep_diary']; c['derived_items'][0]['bins'][0]['lower'] = 1    # closed lower end
        r = ReportEngine(c).score(dict(base, sleep_hours='0.5'))
        self.assertEqual(r['status'], 'blocked'); self.assertIn('outside the configured bins', r['reasons'][0])

    def test_computed_item_configuration_errors(self):
        def d(c, i): return c['derived_items'][i]
        deep = {'const': 1}
        for _ in range(40): deep = {'op': 'abs', 'args': [deep]}
        changes = [lambda c: d(c, 3).update(compute={'item': 'later'}),
                   lambda c: d(c, 0).update(compute={'item': 'c_efficiency'}),                   # defined later
                   lambda c: d(c, 0).update(compute={'op': 'pow', 'args': [{'const': 1}, {'const': 2}]}),
                   lambda c: d(c, 0).update(compute={'op': 'sub', 'args': [{'const': 1}]}),
                   lambda c: d(c, 0).update(compute=deep),
                   lambda c: d(c, 0)['bins'].pop(1),                                                 # gap between bins
                   lambda c: d(c, 0)['bins'][0].update(value=4),                                     # outside 0-3
                   lambda c: d(c, 0).update(id='q1'),
                   lambda c: c['variables'][0].update(type='date')]
        for change in changes:
            c = configurations()['sleep_diary']; change(c)
            with self.assertRaises(ConfigurationError): ReportEngine(c)
        c = configurations()['sleep_diary']
        c['variables'].append(dict(id='group', column='group', type='categorical', codes=['a', 'b']))
        c['derived_items'][0]['compute'] = {'variable': 'group'}
        with self.assertRaises(ConfigurationError): ReportEngine(c)

    # ---------------------------------------------------------------- dynamic report text
    def test_exhaustive_dynamic_text(self):
        e = ReportEngine(configurations()['dynamic_text'])
        header = ['participant_id', 'duration', 'r1', 'r2', 'r3', 'r4', 'r5', 'r6']
        labels = ['Example reason A1', 'Example reason A2', 'Example reason A3', 'Example reason B1', 'Example reason B2', 'Example reason B3']
        lengths = {1: 'less than a month', 2: 'one to six months', 3: 'six to twelve months', 4: 'more than a year'}
        def join(names): return names[0] if len(names) == 1 else ', '.join(names[:-1]) + ' and ' + names[-1]
        cases = [(a, (1, 1, 1)) for a in itertools.product(range(1, 6), repeat=3)] + [((1, 1, 1), b) for b in itertools.product(range(1, 6), repeat=3)]
        for n, (a, b) in enumerate(cases):
            duration = [None, 1, 2, 3, 4][n % 5]
            r = e.score(row_of(header, ['D', duration, *a, *b]))
            texts = {s['id']: s['text'] for s in r['sections']}; decisions = {t['id']: t['decision'] for t in r['section_trace']}
            if duration is None:
                self.assertEqual(decisions['duration_text'], 'withheld_unknown')
            else:
                self.assertEqual(texts['duration_text'], f'You told us this has lasted {lengths[duration]}.')
            if Fraction(sum(a), 3) >= 3:
                top = [labels[i] for i in range(3) if a[i] == max(a)]
                self.assertEqual(texts['top_a'], f'In area A, the reason you rated highest was {join(top)}.')
            else:
                self.assertEqual(decisions['top_a'], 'excluded')
            if Fraction(sum(b), 3) >= 3:
                top = [labels[3 + i] for i in range(3) if b[i] == max(b)]
                if len(top) > 1:
                    self.assertEqual(decisions['top_b'], 'withheld_unknown')
                else:
                    self.assertEqual(texts['top_b'], f'In area B, the reason you rated highest was {top[0]}.')
            else:
                self.assertEqual(decisions['top_b'], 'excluded')
            mean = Fraction(sum(a), 3)
            score = r['scales'][0]['score_display']
            self.assertEqual(Fraction(score), Fraction(half_up(mean * 100), 100))
            band = 'Upper example band' if mean >= 3 else 'Lower example band'
            self.assertEqual(texts['summary'], f'Your area A score was {score} ({band}).')
            COUNTS['dynamic_text_records'] += 1

    def test_first_tie_rule_and_escaping(self):
        c = configurations()['dynamic_text']
        c['sections'][1]['fill']['top']['ties'] = 'first'; c['items'][1]['label'] = '<b>reason</b> & more'
        r = ReportEngine(c).score(row_of(['participant_id', 'duration', 'r1', 'r2', 'r3', 'r4', 'r5', 'r6'], ['E', 2, 5, 3, 5, 1, 1, 1]))
        self.assertEqual(r['sections'][1]['text'], 'In area A, the reason you rated highest was <b>reason</b> & more.')
        h = render_html(r); self.assertIn('&lt;b&gt;reason&lt;/b&gt; &amp; more', h); self.assertNotIn('<b>reason', h)

    def test_dynamic_text_configuration_errors(self):
        def sec(c, i): return c['sections'][i]
        changes = [lambda c: sec(c, 0).update(text='Lasted {{length}} and {{other}}.'),                  # undeclared token
                   lambda c: sec(c, 0).update(text='No token here.'),                                    # unused fill
                   lambda c: sec(c, 1)['fill']['top'].update(ties='random'),
                   lambda c: c['items'][1].pop('label'),                                                # top_items needs labels
                   lambda c: c['items'][0]['response_labels'].pop('4'),                                 # every code labelled
                   lambda c: c['items'][0]['response_labels'].update({'5': 'x'}),
                   lambda c: sec(c, 0).update(fill={'length': {'response_label': 'r1'}}),               # r1 has no labels
                   lambda c: sec(c, 3)['fill'].update(score={'score': 'nope'}),
                   lambda c: c.update(schema_version='1.0')]
        for change in changes:
            c = configurations()['dynamic_text']; change(c)
            with self.assertRaises(ConfigurationError): ReportEngine(c)

    def test_schema_1_0_examples_unchanged(self):
        for name in ('likert', 'weighted_mean', 'knowledge', 'prorated'):
            c = configurations()[name]; self.assertEqual(c['schema_version'], '1.0')
            for field, value in [('variables', []), ('derived_items', [])]:
                bad = deepcopy(c); bad[field] = value
                with self.assertRaises(ConfigurationError): ReportEngine(bad)


if __name__=='__main__':
    import platform,time
    start=time.perf_counter();suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(k) for k in (ScoringChecks,ExtensionChecks)])
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    from qreport import VERSION
    summary=dict(engine_version=VERSION,python=platform.python_version(),tests_run=result.testsRun,
                 failures=len(result.failures),errors=len(result.errors),elapsed_seconds=time.perf_counter()-start,
                 **COUNTS,scope='Invented fixed-form instruments; computational verification only; no human or clinical validation.')
    target=Path(__file__).parent/'validation_results.json';target.write_text(json.dumps(summary,indent=2))
    print(json.dumps(summary,indent=2));raise SystemExit(0 if result.wasSuccessful() else 1)
