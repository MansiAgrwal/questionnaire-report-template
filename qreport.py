"""Configurable fixed-form questionnaire scoring and reporting prototype.

Python 3.10+, standard library for scoring/HTML/JSON; python-docx for Word output.
All scoring and interpretation rules are supplied by the configuration author.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
from copy import deepcopy
from decimal import ROUND_CEILING, ROUND_FLOOR, ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

VERSION = '0.2.0'
SCHEMAS = {'1.0', '1.1'}
METHODS = {'sum', 'mean', 'weighted_sum', 'weighted_mean', 'count'}
COMPARISONS = {'eq', 'ne', 'lt', 'le', 'gt', 'ge'}
IDENTIFIER = re.compile(r'^[A-Za-z][A-Za-z0-9_]*$')
TOKENS = {'score', 'band', 'answered', 'items'}
CLOCK = re.compile(r'^([01]?\d|2[0-3]):([0-5]\d)$')
# operator: (minimum arguments, maximum arguments or None for unlimited)
EXPRESSION_OPS = {'add': (2, None), 'sub': (2, 2), 'mul': (2, None), 'div': (2, 2), 'min': (2, None),
                  'max': (2, None), 'abs': (1, 1), 'minutes_between': (2, 2)}
MAX_EXPRESSION_DEPTH = 32

# Fields added in schema 1.1; a 1.0 configuration may not use them.
NEW_TOP = {'variables', 'derived_items'}
NEW_ITEM = {'label', 'response_labels'}
NEW_SCALE = {'criterion', 'item_criteria', 'lookup'}
NEW_SECTION = {'fill'}


class ConfigurationError(ValueError):
    pass


class DataError(ValueError):
    pass


def number(value):
    if isinstance(value, bool):
        raise ValueError('Boolean values are not numeric scores.')
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f'Not a numeric score: {value!r}') from None
    if not result.is_finite():
        raise ValueError('Scores must be finite.')
    return result


def key(value):
    if value is None:
        return ''
    text = str(value).strip()
    try:
        x = number(text)
        return format(x.normalize(), 'f')
    except ValueError:
        return text


def exact(value):
    return format(value.normalize(), 'f') if value else '0'


def shown(value, decimals):
    """Participant-facing score: rounded half up to `decimals` places, trailing zeros removed."""
    return exact(value.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP))


def clock(value):
    """Minutes after midnight for a 24-hour H:MM or HH:MM time."""
    match = CLOCK.fullmatch(str(value).strip())
    if not match:
        raise ValueError(f'Not a 24-hour clock time (HH:MM): {value!r}')
    return Decimal(int(match[1]) * 60 + int(match[2]))


def rounded(value, rounding, low=None):
    """Raw score as used for a table lookup. For range limits under rounding 'none', low/high pick ceiling/floor."""
    if rounding == 'half_up':
        return value.quantize(Decimal(1), rounding=ROUND_HALF_UP)
    if rounding == 'up':
        return value.to_integral_value(rounding=ROUND_CEILING)
    if low is None:
        return value
    return value.to_integral_value(rounding=ROUND_CEILING if low else ROUND_FLOOR)


def compare(value, op, rhs):
    return {'eq': value == rhs, 'ne': value != rhs, 'lt': value < rhs, 'le': value <= rhs,
            'gt': value > rhs, 'ge': value >= rhs}[op]


def join_labels(labels):
    return labels[0] if len(labels) == 1 else ', '.join(labels[:-1]) + ' and ' + labels[-1]


def known_keys(obj, allowed, location):
    if not isinstance(obj, dict):
        raise ConfigurationError(f'{location}: expected an object.')
    extra = set(obj) - set(allowed)
    if extra:
        raise ConfigurationError(f'{location}: unknown fields {sorted(extra)}')


def required_text(obj, field, location):
    if not isinstance(obj.get(field), str) or not obj[field].strip():
        raise ConfigurationError(f'{location}: {field} must be a nonempty string.')


def identifier(value, location):
    if not isinstance(value, str) or not IDENTIFIER.fullmatch(value):
        raise ConfigurationError(f'{location}: use a letter followed by letters, digits or underscores.')


def check_template(text):
    if not isinstance(text, str):
        raise ConfigurationError('Band text must be a string.')
    found = re.findall(r'\{\{([a-z_]+)\}\}', text)
    rest = re.sub(r'\{\{([a-z_]+)\}\}', '', text)
    if set(found) - TOKENS or '{{' in rest or '}}' in rest:
        raise ConfigurationError('Templates allow only {{score}}, {{band}}, {{answered}} and {{items}}.')


def load_config(path):
    def unique_pairs(pairs):
        out = {}
        for k, v in pairs:
            if k in out:
                raise ConfigurationError(f'Duplicate JSON key: {k}')
            out[k] = v
        return out
    return json.loads(Path(path).read_text(encoding='utf-8'), object_pairs_hook=unique_pairs)


def check_intervals(rows, lower, upper, location, value_field, open_ends=False, inner_inclusive=False):
    """Contiguous ascending intervals covering [lower, upper] (or the whole line when open_ends).

    Lower bounds are inclusive. Without inner_inclusive only the final interval includes its upper bound; with it,
    any interval may set upper_inclusive, and the next interval then excludes that shared boundary.
    """
    if not isinstance(rows, list) or not rows:
        raise ConfigurationError(f'{location}: a nonempty list is required.')
    out = []
    for j, row in enumerate(rows):
        known_keys(row, {'lower', 'upper', 'upper_inclusive', value_field}, location)
        if value_field not in row:
            raise ConfigurationError(f'{location}: each interval needs {value_field}.')
        first, last = j == 0, j == len(rows) - 1
        lo = None if open_ends and first and row.get('lower') is None else number(row['lower'])
        hi = None if open_ends and last and row.get('upper') is None else number(row['upper'])
        inclusive = row.get('upper_inclusive', False)
        if not isinstance(inclusive, bool):
            raise ConfigurationError(f'{location}: upper_inclusive must be boolean.')
        if lo is not None and hi is not None and lo >= hi:
            raise ConfigurationError(f'{location}: each lower bound must be below its upper bound.')
        if not inner_inclusive and inclusive != (last and hi is not None):
            raise ConfigurationError(f'{location}: only the final interval includes its upper boundary.')
        if hi is None and inclusive:
            raise ConfigurationError(f'{location}: an open upper end cannot be inclusive.')
        if out and out[-1]['upper'] != lo:
            raise ConfigurationError(f'{location}: intervals must be contiguous, without gaps or overlap.')
        out.append(dict(lower=lo, upper=hi, inclusive=inclusive, value=row[value_field],
                        lower_exclusive=bool(out and out[-1]['inclusive'])))
    if not open_ends and (out[0]['lower'] != lower or out[-1]['upper'] != upper):
        raise ConfigurationError(f'{location}: intervals must cover the range {exact(lower)} to {exact(upper)}.')
    if not open_ends and not out[-1]['inclusive']:
        raise ConfigurationError(f'{location}: the final interval must include its upper boundary.')
    return out


def find_interval(intervals, value):
    for row in intervals:
        above = row['lower'] is None or (value > row['lower'] if row['lower_exclusive'] else value >= row['lower'])
        below = row['upper'] is None or value < row['upper'] or (row['inclusive'] and value == row['upper'])
        if above and below:
            return row
    return None


class ReportEngine:
    def __init__(self, config):
        self.config = deepcopy(config)
        self.items, self.variables, self.scales = {}, {}, []
        try:
            self._compile()
        except ConfigurationError:
            raise
        except (ValueError, TypeError, KeyError, AttributeError) as exc:
            raise ConfigurationError(str(exc)) from exc
        canonical = json.dumps(self.config, sort_keys=True, ensure_ascii=False, separators=(',', ':'))
        self.config_sha256 = hashlib.sha256(canonical.encode()).hexdigest()

    # ----------------------------------------------------------------------------------------------- compile
    def _compile(self):
        c = self.config
        if not isinstance(c, dict) or c.get('schema_version') not in SCHEMAS:
            raise ConfigurationError(f'schema_version must be one of {sorted(SCHEMAS)}.')
        new = c['schema_version'] != '1.0'
        top = {'schema_version', 'questionnaire_id', 'title', 'instrument_version', 'scoring_reference', 'id_column',
               'introduction', 'missing_codes', 'eligibility', 'items', 'scales', 'sections'}
        known_keys(c, top | (NEW_TOP if new else set()), 'configuration')
        for f in ('questionnaire_id', 'title', 'instrument_version', 'scoring_reference', 'id_column'):
            required_text(c, f, 'configuration')
        identifier(c['questionnaire_id'], 'questionnaire_id')
        if not isinstance(c.get('introduction', ''), str):
            raise ConfigurationError('introduction must be text.')
        missing = c.get('missing_codes', ['', None, 'NA'])
        if not isinstance(missing, list):
            raise ConfigurationError('missing_codes must be a list.')
        self.missing = {key(x) for x in missing}
        if not isinstance(c.get('items'), list) or not c['items']:
            raise ConfigurationError('At least one item is required.')
        self.used_columns = {c['id_column']}
        for item in c['items']:
            self._compile_item(item, new)
        if not isinstance(c.get('variables', []), list):
            raise ConfigurationError('variables must be a list.')
        for v in c.get('variables', []):
            self._compile_variable(v)
        derived = c.get('derived_items', [])
        if not isinstance(derived, list):
            raise ConfigurationError('derived_items must be a list.')
        for d in derived:
            self._compile_derived(d)
        self.eligibility = c.get('eligibility', [])
        if not isinstance(self.eligibility, list):
            raise ConfigurationError('eligibility must be a list.')
        for e in self.eligibility:
            known_keys(e, {'column', 'equals', 'reason'}, 'eligibility')
            for f in ('column', 'reason'):
                required_text(e, f, 'eligibility')
            if 'equals' not in e:
                raise ConfigurationError('Eligibility requires equals.')
        if not isinstance(c.get('scales'), list) or not c['scales']:
            raise ConfigurationError('At least one scale is required.')
        self.scale_ids = set()
        for s in c['scales']:
            self._compile_scale(s, new)
        self.sections = c.get('sections', [])
        if not isinstance(self.sections, list):
            raise ConfigurationError('sections must be a list.')
        seen = set()
        for section in self.sections:
            known_keys(section, {'id', 'title', 'text', 'source', 'when'} | (NEW_SECTION if new else set()), 'section')
            identifier(section.get('id'), 'section id')
            if section['id'] in seen:
                raise ConfigurationError('Duplicate section id.')
            seen.add(section['id'])
            for f in ('title', 'text', 'source'):
                required_text(section, f, 'section')
            fill = section.get('fill', {})
            if not isinstance(fill, dict):
                raise ConfigurationError('fill must be an object.')
            found = re.findall(r'\{\{([A-Za-z][A-Za-z0-9_]*)\}\}', section['text'])
            rest = re.sub(r'\{\{([A-Za-z][A-Za-z0-9_]*)\}\}', '', section['text'])
            if '{{' in rest or '}}' in rest or set(found) != set(fill):
                raise ConfigurationError(f'Section {section["id"]}: every {{{{token}}}} in the text must be declared in '
                                         'fill, and every fill entry must be used. Without fill, section text is literal.')
            for token, spec in fill.items():
                identifier(token, 'fill token')
                self._check_fill(spec, f'section {section["id"]}, fill {token}')
            if 'when' in section:
                self._check_condition(section['when'], 'section condition')

    def _compile_item(self, item, new):
        known_keys(item, {'id', 'column', 'min', 'max', 'reverse', 'allow_decimal', 'recode'} | (NEW_ITEM if new else set()), 'item')
        identifier(item.get('id'), 'item id')
        required_text(item, 'column', 'item')
        iid = item['id']
        if iid in self.items or item['column'] in self.used_columns:
            raise ConfigurationError('Item ids and item columns must be unique and distinct from id_column.')
        lo, hi = number(item['min']), number(item['max'])
        if lo > hi:
            raise ConfigurationError(f'{iid}: min exceeds max.')
        for f in ('reverse', 'allow_decimal'):
            if f in item and not isinstance(item[f], bool):
                raise ConfigurationError(f'{iid}: {f} must be boolean.')
        decimal_ok = item.get('allow_decimal', False)
        recode = item.get('recode')
        if recode is not None:
            if not isinstance(recode, dict) or not recode:
                raise ConfigurationError(f'{iid}: recode must be a nonempty object.')
            mapped = {key(k): number(v) for k, v in recode.items()}
            if len(mapped) != len(recode) or self.missing.intersection(mapped):
                raise ConfigurationError(f'{iid}: duplicate or missing response code in recode.')
            if any(v < lo or v > hi for v in mapped.values()):
                raise ConfigurationError(f'{iid}: recoded value outside item bounds.')
            if not decimal_ok and any(v != v.to_integral_value() for v in mapped.values()):
                raise ConfigurationError(f'{iid}: fractional recoding requires allow_decimal.')
        else:
            mapped = None
            self._missing_overlap(lo, hi, decimal_ok, iid)
        if 'label' in item:
            required_text(item, 'label', f'item {iid}')
        labels = None
        if 'response_labels' in item:
            raw = item['response_labels']
            if not isinstance(raw, dict) or not raw or not all(isinstance(v, str) and v.strip() for v in raw.values()):
                raise ConfigurationError(f'{iid}: response_labels must map response codes to nonempty text.')
            labels = {key(k): v for k, v in raw.items()}
            if len(labels) != len(raw):
                raise ConfigurationError(f'{iid}: duplicate response code in response_labels.')
            if mapped is not None:
                codes = set(mapped)
            elif not decimal_ok:
                if hi - lo > 1000 or lo != lo.to_integral_value() or hi != hi.to_integral_value():
                    raise ConfigurationError(f'{iid}: response_labels need a recode map or a short integer range.')
                codes = {key(Decimal(v)) for v in range(int(lo), int(hi) + 1)}
            else:
                raise ConfigurationError(f'{iid}: response_labels cannot cover a continuous response.')
            if set(labels) != codes:
                raise ConfigurationError(f'{iid}: response_labels must label every valid response code exactly once.')
        self.items[iid] = dict(source=item, lo=lo, hi=hi, recode=mapped, labels=labels, derived=False)
        self.used_columns.add(item['column'])

    def _missing_overlap(self, lo, hi, decimal_ok, name):
        for m in self.missing:
            try:
                x = number(m)
            except ValueError:
                continue
            if lo <= x <= hi and (decimal_ok or x == x.to_integral_value()):
                raise ConfigurationError(f'{name}: missing code {m} overlaps a valid response.')

    def _compile_variable(self, v):
        known_keys(v, {'id', 'column', 'type', 'min', 'max', 'allow_decimal', 'codes'}, 'variable')
        identifier(v.get('id'), 'variable id')
        required_text(v, 'column', 'variable')
        if v['id'] in self.variables or v['column'] in self.used_columns:
            raise ConfigurationError('Variable ids must be unique, and variable columns distinct from item and id columns.')
        kind = v.get('type')
        spec = dict(source=v, type=kind)
        if kind == 'numeric':
            spec['lo'], spec['hi'] = number(v['min']), number(v['max'])
            if spec['lo'] > spec['hi']:
                raise ConfigurationError(f'{v["id"]}: min exceeds max.')
            if not isinstance(v.get('allow_decimal', False), bool):
                raise ConfigurationError(f'{v["id"]}: allow_decimal must be boolean.')
            self._missing_overlap(spec['lo'], spec['hi'], v.get('allow_decimal', False), v['id'])
            if 'codes' in v:
                raise ConfigurationError(f'{v["id"]}: codes apply only to categorical variables.')
        elif kind == 'categorical':
            codes = v.get('codes')
            if not isinstance(codes, list) or not codes:
                raise ConfigurationError(f'{v["id"]}: a categorical variable needs a nonempty codes list.')
            spec['codes'] = {key(x) for x in codes}
            if len(spec['codes']) != len(codes) or spec['codes'] & self.missing:
                raise ConfigurationError(f'{v["id"]}: codes must be unique and distinct from missing codes.')
            if set(v) & {'min', 'max', 'allow_decimal'}:
                raise ConfigurationError(f'{v["id"]}: bounds apply only to numeric variables.')
        elif kind == 'time':
            if set(v) & {'min', 'max', 'allow_decimal', 'codes'}:
                raise ConfigurationError(f'{v["id"]}: a time variable takes no bounds or codes.')
        else:
            raise ConfigurationError(f'{v.get("id")}: variable type must be numeric, categorical or time.')
        self.variables[v['id']] = spec
        self.used_columns.add(v['column'])

    def _compile_derived(self, d):
        known_keys(d, {'id', 'label', 'compute', 'bins', 'min', 'max', 'allow_decimal'}, 'derived item')
        identifier(d.get('id'), 'derived item id')
        if d['id'] in self.items:
            raise ConfigurationError('Derived item ids must differ from item ids.')
        lo, hi = number(d['min']), number(d['max'])
        if lo > hi:
            raise ConfigurationError(f'{d["id"]}: min exceeds max.')
        decimal_ok = d.get('allow_decimal', False)
        if not isinstance(decimal_ok, bool):
            raise ConfigurationError(f'{d["id"]}: allow_decimal must be boolean.')
        if 'label' in d:
            required_text(d, 'label', f'derived item {d["id"]}')
        self._check_expression(d.get('compute'), d['id'], 0)
        bins = None
        if 'bins' in d:
            bins = check_intervals(d['bins'], None, None, f'{d["id"]} bins', 'value', open_ends=True, inner_inclusive=True)
            for b in bins:
                b['value'] = number(b['value'])
                if not lo <= b['value'] <= hi or (not decimal_ok and b['value'] != b['value'].to_integral_value()):
                    raise ConfigurationError(f'{d["id"]}: each bin value must be a valid score for this item.')
        self.items[d['id']] = dict(source=d, lo=lo, hi=hi, recode=None, labels=None, derived=True, bins=bins)

    def _check_expression(self, node, owner, depth):
        if depth > MAX_EXPRESSION_DEPTH:
            raise ConfigurationError(f'{owner}: expression is nested too deeply.')
        if not isinstance(node, dict) or len(node) not in (1, 2):
            raise ConfigurationError(f'{owner}: each expression node is a const, item, variable or op object.')
        if set(node) == {'const'}:
            number(node['const'])
        elif set(node) == {'item'}:
            if node['item'] not in self.items:
                raise ConfigurationError(f'{owner}: expression refers to an undefined or later item {node["item"]!r}.')
        elif set(node) == {'variable'}:
            v = self.variables.get(node['variable'])
            if v is None or v['type'] == 'categorical':
                raise ConfigurationError(f'{owner}: expressions use numeric or time variables only.')
        elif set(node) == {'op', 'args'}:
            if node['op'] not in EXPRESSION_OPS:
                raise ConfigurationError(f'{owner}: unknown operation {node["op"]!r}.')
            low, high = EXPRESSION_OPS[node['op']]
            args = node['args']
            if not isinstance(args, list) or len(args) < low or (high is not None and len(args) > high):
                raise ConfigurationError(f'{owner}: wrong number of arguments for {node["op"]}.')
            for a in args:
                self._check_expression(a, owner, depth + 1)
        else:
            raise ConfigurationError(f'{owner}: each expression node is a const, item, variable or op object.')

    def _compile_scale(self, s, new):
        allowed = {'id', 'label', 'items', 'method', 'weights', 'missing_policy', 'multiplier', 'offset', 'bands',
                   'interpretation_reference', 'display_decimals'} | (NEW_SCALE if new else set())
        known_keys(s, allowed, 'scale')
        identifier(s.get('id'), 'scale id')
        required_text(s, 'label', 'scale')
        if s['id'] in self.scale_ids:
            raise ConfigurationError('Duplicate scale id.')
        if not isinstance(s.get('items'), list) or not s['items'] or len(s['items']) != len(set(s['items'])):
            raise ConfigurationError('Scale items must be a nonempty unique list.')
        if set(s['items']) - set(self.items):
            raise ConfigurationError('Scale references an undefined item.')
        method = s.get('method')
        if method not in METHODS:
            raise ConfigurationError(f'Unknown scoring method: {method}')
        weights = s.get('weights', {})
        if not isinstance(weights, dict) or set(weights) - set(s['items']):
            raise ConfigurationError('Weights reference an undefined scale item.')
        if weights and not method.startswith('weighted_'):
            raise ConfigurationError('Use a weighted method when supplying weights.')
        w = {i: number(weights.get(i, 1)) for i in s['items']}
        if any(x <= 0 for x in w.values()):
            raise ConfigurationError('Weights must be positive.')
        mp = s.get('missing_policy', {'method': 'complete'})
        known_keys(mp, {'method', 'minimum_answered', 'justification'}, 'missing_policy')
        if mp.get('method') not in {'complete', 'prorate'}:
            raise ConfigurationError('Missing policy must be complete or prorate.')
        n = len(s['items'])
        minimum = n
        if mp['method'] == 'prorate':
            if method == 'count':
                raise ConfigurationError('Count scales require complete responses; proration does not apply to counts.')
            minimum = mp.get('minimum_answered')
            if isinstance(minimum, bool) or not isinstance(minimum, int) or not 1 <= minimum <= n:
                raise ConfigurationError('Proration requires a valid minimum_answered.')
            required_text(mp, 'justification', 'proration')
            bounds = {(self.items[i]['lo'], self.items[i]['hi']) for i in s['items']}
            if len(bounds) != 1:
                raise ConfigurationError('This prototype prorates only items with identical scored bounds.')
        elif 'minimum_answered' in mp or 'justification' in mp:
            raise ConfigurationError('Complete scoring does not accept proration settings.')
        criteria = None
        if method == 'count':
            if 'criterion' not in s:
                raise ConfigurationError('A count scale needs a criterion.')
            overrides = s.get('item_criteria', {})
            if not isinstance(overrides, dict) or set(overrides) - set(s['items']):
                raise ConfigurationError('item_criteria must name items of this scale.')
            criteria = {}
            for i in s['items']:
                crit = overrides.get(i, s['criterion'])
                known_keys(crit, {'op', 'value'}, 'count criterion')
                if crit.get('op') not in COMPARISONS:
                    raise ConfigurationError('Unknown comparison operator in count criterion.')
                criteria[i] = (crit['op'], number(crit['value']))
        elif 'criterion' in s or 'item_criteria' in s:
            raise ConfigurationError('criterion and item_criteria apply only to count scales.')
        decimals = s.get('display_decimals', 2)
        if isinstance(decimals, bool) or not isinstance(decimals, int) or not 0 <= decimals <= 6:
            raise ConfigurationError('display_decimals must be an integer from 0 to 6.')
        mult, offset = number(s.get('multiplier', 1)), number(s.get('offset', 0))
        if mult <= 0:
            raise ConfigurationError('multiplier must be positive.')
        if method == 'count':
            raw_lower, raw_upper = Decimal(0), Decimal(n)
        else:
            raw_lower = sum(self.items[i]['lo'] * w[i] for i in s['items'])
            raw_upper = sum(self.items[i]['hi'] * w[i] for i in s['items'])
            if method.endswith('mean'):
                raw_lower /= sum(w.values()); raw_upper /= sum(w.values())
        lookup = None
        if 'lookup' in s:
            if 'multiplier' in s or 'offset' in s:
                raise ConfigurationError('Use either a lookup table or multiplier/offset, not both.')
            lookup = self._compile_lookup(s, raw_lower, raw_upper, w, method, mp['method'])
            lower, upper = lookup['lower'], lookup['upper']
        else:
            lower, upper = raw_lower * mult + offset, raw_upper * mult + offset
        bands = []
        source_bands = s.get('bands', [])
        if not isinstance(source_bands, list):
            raise ConfigurationError('bands must be a list.')
        if source_bands:
            required_text(s, 'interpretation_reference', 'scale with bands')
        seen_bands = set()
        for j, b in enumerate(source_bands):
            known_keys(b, {'id', 'label', 'lower', 'upper', 'upper_inclusive', 'text', 'hold'}, 'band')
            identifier(b.get('id'), 'band id'); required_text(b, 'label', 'band')
            if b['id'] in seen_bands:
                raise ConfigurationError('Band ids must be unique within a scale.')
            seen_bands.add(b['id'])
            bl, bu = number(b['lower']), number(b['upper'])
            last = j == len(source_bands) - 1
            if not isinstance(b.get('upper_inclusive', False), bool) or not isinstance(b.get('hold', False), bool):
                raise ConfigurationError('Band flags must be boolean.')
            if bl >= bu or b.get('upper_inclusive', False) != last:
                raise ConfigurationError('Only the final band must include its upper boundary.')
            if j == 0 and bl != lower or bands and bands[-1]['upper'] != bl:
                raise ConfigurationError('Bands must cover the score range without gaps or overlap, in ascending order.')
            check_template(b.get('text', ''))
            bands.append(dict(source=b, lower=bl, upper=bu))
        if bands and bands[-1]['upper'] != upper:
            raise ConfigurationError('Bands do not cover the maximum score.')
        self.scale_ids.add(s['id'])
        self.scales.append(dict(source=s, weights=w, minimum=minimum, mult=mult, offset=offset, decimals=decimals,
                                bands=bands, lower=lower, upper=upper, missing_policy=mp['method'],
                                criteria=criteria, lookup=lookup))

    def _compile_lookup(self, s, raw_lower, raw_upper, w, method, policy):
        lk = s['lookup']
        known_keys(lk, {'reference', 'round_raw', 'tables'}, 'lookup')
        required_text(lk, 'reference', 'lookup')
        rounding = lk.get('round_raw', 'none')
        if rounding not in {'none', 'half_up', 'up'}:
            raise ConfigurationError('round_raw must be none, half_up or up.')
        tables = lk.get('tables')
        if not isinstance(tables, list) or not tables:
            raise ConfigurationError('lookup needs a nonempty list of tables.')
        integral_raw = rounding != 'none' or (
            method in {'sum', 'weighted_sum', 'count'} and policy == 'complete'
            and all(x == x.to_integral_value() for x in w.values())
            and all(not self.items[i]['source'].get('allow_decimal', False)
                    and all(x == x.to_integral_value() for x in (self.items[i]['lo'], self.items[i]['hi'])) for i in s['items']))
        compiled, ids, outputs = [], set(), []
        for t in tables:
            known_keys(t, {'id', 'group', 'type', 'rows'}, 'lookup table')
            identifier(t.get('id'), 'lookup table id')
            if t['id'] in ids:
                raise ConfigurationError('Lookup table ids must be unique within a scale.')
            ids.add(t['id'])
            if 'group' in t:
                self._check_condition(t['group'], 'lookup group', variables_only=True)
            if t.get('type') == 'exact':
                if not integral_raw:
                    raise ConfigurationError('An exact table needs whole-number raw scores; set round_raw or use an interval table.')
                rows = t.get('rows')
                if not isinstance(rows, list) or not rows:
                    raise ConfigurationError('An exact table needs rows.')
                table = {}
                for row in rows:
                    known_keys(row, {'raw', 'score'}, 'exact table row')
                    r = number(row['raw'])
                    if r != r.to_integral_value() or key(r) in table:
                        raise ConfigurationError('Exact table raw values must be unique whole numbers.')
                    table[key(r)] = number(row['score'])
                first, last = (rounded(raw_lower, rounding, low=True), rounded(raw_upper, rounding, low=False))
                expected = {key(Decimal(v)) for v in range(int(first), int(last) + 1)}
                if set(table) != expected:
                    raise ConfigurationError(f'Exact table {t["id"]} must list every whole raw score from {exact(first)} to '
                                             f'{exact(last)} exactly once.')
                outputs.extend(table.values())
                compiled.append(dict(source=t, type='exact', table=table))
            elif t.get('type') == 'interval':
                rows = check_intervals(t.get('rows'), raw_lower, raw_upper, f'lookup table {t["id"]}', 'score')
                for row in rows:
                    row['value'] = number(row['value'])
                outputs.extend(row['value'] for row in rows)
                compiled.append(dict(source=t, type='interval', rows=rows))
            else:
                raise ConfigurationError('Lookup table type must be exact or interval.')
        return dict(reference=lk['reference'], rounding=rounding, tables=compiled, lower=min(outputs), upper=max(outputs))

    def _check_condition(self, node, location, variables_only=False):
        if not isinstance(node, dict):
            raise ConfigurationError(f'{location}: a condition must be an object.')
        if set(node) in ({'all'}, {'any'}):
            nodes = next(iter(node.values()))
            if not isinstance(nodes, list) or not nodes:
                raise ConfigurationError('all/any conditions require a nonempty list.')
            for n in nodes:
                self._check_condition(n, location, variables_only)
            return
        if set(node) == {'not'}:
            self._check_condition(node['not'], location, variables_only)
            return
        known_keys(node, {'scale', 'item', 'variable', 'op', 'value'}, 'condition')
        targets = [k for k in ('scale', 'item', 'variable') if k in node]
        if len(targets) != 1:
            raise ConfigurationError('A condition names exactly one scale, item or variable.')
        if variables_only and targets != ['variable']:
            raise ConfigurationError(f'{location}: only variables can be used here.')
        if node.get('op') not in COMPARISONS:
            raise ConfigurationError('Unknown comparison operator.')
        if 'scale' in node and node['scale'] not in self.scale_ids:
            raise ConfigurationError('Condition references an undefined scale/item.')
        if 'item' in node and node['item'] not in self.items:
            raise ConfigurationError('Condition references an undefined scale/item.')
        if 'variable' in node:
            v = self.variables.get(node['variable'])
            if v is None:
                raise ConfigurationError('Condition references an undefined variable.')
            if v['type'] == 'categorical':
                if node['op'] not in {'eq', 'ne'} or key(node['value']) not in v['codes']:
                    raise ConfigurationError('Categorical conditions use eq or ne with one of the declared codes.')
                return
            if v['type'] == 'time' and isinstance(node['value'], str):
                clock(node['value'])
                return
        number(node['value'])

    def _check_fill(self, spec, location):
        if not isinstance(spec, dict):
            raise ConfigurationError(f'{location}: expected an object.')
        if set(spec) in ({'top_items', 'ties'}, {'top_items'}):
            if spec['top_items'] not in self.scale_ids:
                raise ConfigurationError(f'{location}: top_items names an undefined scale.')
            if spec.get('ties', 'all') not in {'all', 'first', 'withhold'}:
                raise ConfigurationError(f'{location}: ties must be all, first or withhold.')
            scale = next(x for x in self.scales if x['source']['id'] == spec['top_items'])
            if any('label' not in self.items[i]['source'] for i in scale['source']['items']):
                raise ConfigurationError(f'{location}: every item of the scale needs a label.')
        elif set(spec) == {'response_label'}:
            item = self.items.get(spec['response_label'])
            if item is None or item['labels'] is None:
                raise ConfigurationError(f'{location}: response_label needs an item with response_labels.')
        elif set(spec) in ({'score'}, {'band'}):
            if next(iter(spec.values())) not in self.scale_ids:
                raise ConfigurationError(f'{location}: names an undefined scale.')
        else:
            raise ConfigurationError(f'{location}: use top_items (with optional ties), response_label, score or band.')

    # ----------------------------------------------------------------------------------------------- score
    def score(self, row):
        c = self.config
        needed = ({c['id_column']} | {i['source']['column'] for i in self.items.values() if not i['derived']}
                  | {v['source']['column'] for v in self.variables.values()} | {e['column'] for e in self.eligibility})
        absent = needed - set(row)
        if absent:
            raise DataError(f'Missing input columns: {sorted(absent)}')
        pid = '' if row[c['id_column']] is None else str(row[c['id_column']]).strip()
        if not pid or len(pid) > 128 or any(ord(x) < 32 for x in pid):
            raise DataError('Participant id must be printable, nonempty and at most 128 characters.')
        reasons = [e['reason'] for e in self.eligibility if key(row[e['column']]) != key(e['equals'])]
        trace, variable_trace = {}, {}
        for iid, item in self.items.items():
            if item['derived']:
                continue
            spec = item['source']; raw = row[spec['column']]; token = key(raw)
            t = dict(column=spec['column'], raw=None if raw is None else str(raw), status='missing')
            if token not in self.missing:
                try:
                    if len(token) > 256:
                        raise ValueError('Response token too long.')
                    value = item['recode'][token] if item['recode'] is not None else number(token)
                    if not item['lo'] <= value <= item['hi']:
                        raise ValueError('Response outside configured bounds.')
                    if not spec.get('allow_decimal', False) and value != value.to_integral_value():
                        raise ValueError('Fractional response is not allowed for this item.')
                    scored = item['lo'] + item['hi'] - value if spec.get('reverse', False) else value
                    t.update(status='valid', code=token, mapped=exact(value), scored=exact(scored), reversed=spec.get('reverse', False))
                except (ValueError, KeyError) as exc:
                    t.update(status='invalid', reason='Unknown response code.' if isinstance(exc, KeyError) else str(exc))
                    reasons.append(f'Invalid response in {iid}: {t["reason"]}')
            trace[iid] = t
        for vid, v in self.variables.items():
            raw = row[v['source']['column']]; token = key(raw)
            t = dict(column=v['source']['column'], raw=None if raw is None else str(raw), status='missing')
            if token not in self.missing:
                try:
                    if len(token) > 256:
                        raise ValueError('Value too long.')
                    if v['type'] == 'categorical':
                        if token not in v['codes']:
                            raise ValueError('Unknown category code.')
                        value = token
                    elif v['type'] == 'time':
                        value = exact(clock(token))
                    else:
                        x = number(token)
                        if not v['lo'] <= x <= v['hi']:
                            raise ValueError('Value outside configured bounds.')
                        if not v['source'].get('allow_decimal', False) and x != x.to_integral_value():
                            raise ValueError('Fractional value is not allowed for this variable.')
                        value = exact(x)
                    t.update(status='valid', value=value)
                except ValueError as exc:
                    t.update(status='invalid', reason=str(exc))
                    reasons.append(f'Invalid value in {vid}: {exc}')
            variable_trace[vid] = t
        for iid, item in self.items.items():
            if not item['derived']:
                continue
            t = dict(derived=True, status='missing')
            try:
                value = self._evaluate_expression(item['source']['compute'], trace, variable_trace)
                if value is not None:
                    t['computed'] = exact(value)
                    if item['bins'] is not None:
                        hit = find_interval(item['bins'], value)
                        if hit is None:
                            raise ValueError('Computed value outside the configured bins.')
                        value = hit['value']
                    elif not item['lo'] <= value <= item['hi']:
                        raise ValueError('Computed value outside configured bounds.')
                    elif not item['source'].get('allow_decimal', False) and value != value.to_integral_value():
                        raise ValueError('Computed value is fractional but allow_decimal is not set.')
                    t.update(status='valid', scored=exact(value))
            except ValueError as exc:
                t.update(status='invalid', reason=str(exc))
                reasons.append(f'Invalid computed item {iid}: {exc}')
            trace[iid] = t
        results, hold = [], False
        for scale in self.scales:
            r, held = self._score_scale(scale, trace, variable_trace)
            results.append(r); hold = hold or held
        scored = sum(r['status'] == 'scored' for r in results)
        status = 'blocked' if reasons else 'held_for_review' if hold else 'unscored' if scored == 0 else 'partial' if scored < len(results) else 'ready'
        by_id = {r['id']: r for r in results}
        sections, section_trace = [], []
        for s in self.sections:
            decision = self._evaluate(s['when'], trace, variable_trace, by_id) if 'when' in s else True
            entry = dict(id=s['id'], source=s['source'])
            text = s['text']
            if decision is True and s.get('fill'):
                values = {}
                for token, spec in s['fill'].items():
                    values[token] = self._fill_value(spec, trace, by_id)
                entry['fill'] = {k: v for k, v in values.items()}
                if any(v is None for v in values.values()):
                    decision = None
                else:
                    text = re.sub(r'\{\{([A-Za-z][A-Za-z0-9_]*)\}\}', lambda m: values[m[1]], text)
            entry['decision'] = 'included' if decision is True else 'excluded' if decision is False else 'withheld_unknown'
            section_trace.append(entry)
            if decision is True:
                sections.append(dict(id=s['id'], title=s['title'], text=text, source=s['source']))
        return dict(engine_version=VERSION, questionnaire_id=c['questionnaire_id'], instrument_version=c['instrument_version'],
                    config_sha256=self.config_sha256, participant_id=pid, title=c['title'],
                    introduction=c.get('introduction', ''), scoring_reference=c['scoring_reference'],
                    status=status, reasons=reasons, scales=results, item_trace=trace, variable_trace=variable_trace,
                    sections=sections, section_trace=section_trace)

    def _score_scale(self, scale, trace, variable_trace):
        s = scale['source']; use = s['items']; w = scale['weights']
        good = [i for i in use if trace[i]['status'] == 'valid']
        invalid = [i for i in use if trace[i]['status'] == 'invalid']
        r = dict(id=s['id'], label=s['label'], method=s['method'], item_ids=use,
                 answered=len(good), expected=len(use), missing=[i for i in use if trace[i]['status'] == 'missing'],
                 invalid=invalid, status='unscored', score=None, score_exact=None, score_display=None, band=None, text='', prorated=False)
        if invalid:
            r['reason'] = 'Invalid item response; scale withheld.'
            return r, False
        if len(good) < scale['minimum']:
            r['reason'] = 'Too few answered items for the configured scoring rule.'
            return r, False
        if s['method'] == 'count':
            met = [i for i in use if compare(number(trace[i]['scored']), *scale['criteria'][i])]
            raw_score = Decimal(len(met))
            r['items_meeting_criterion'] = met
        else:
            numerator = sum(number(trace[i]['scored']) * w[i] for i in good)
            denominator = sum(w[i] for i in good)
            raw_score = numerator / denominator if s['method'].endswith('mean') else numerator
            if len(good) < len(use):
                if not s['method'].endswith('mean'):
                    raw_score *= sum(w.values()) / denominator
                r['prorated'] = True
                r['proration_justification'] = s['missing_policy']['justification']
        r['raw_score_exact'] = exact(raw_score)
        lk = scale['lookup']
        if lk is None:
            value = raw_score * scale['mult'] + scale['offset']
        else:
            used = raw_score
            used = rounded(raw_score, lk['rounding'])
            table = None
            for t in lk['tables']:
                g = self._evaluate(t['source']['group'], trace, variable_trace, {}) if 'group' in t['source'] else True
                if g is None:
                    r['reason'] = 'The conversion table could not be chosen because a grouping variable is missing.'
                    return r, False
                if g:
                    table = t
                    break
            if table is None:
                r['reason'] = 'No conversion table applies to this participant.'
                return r, False
            if table['type'] == 'exact':
                if key(used) not in table['table']:
                    raise DataError('A raw score is missing from an exact conversion table.')
                value = table['table'][key(used)]
            else:
                value = find_interval(table['rows'], used)['value']
            r['lookup'] = dict(table=table['source']['id'], raw_used=exact(used), reference=lk['reference'])
        r.update(status='scored', score=float(value), score_exact=exact(value), score_display=shown(value, scale['decimals']))
        held = False
        for b in scale['bands']:
            bs = b['source']
            if value >= b['lower'] and (value < b['upper'] or bs['upper_inclusive'] and value == b['upper']):
                r['band'] = dict(id=bs['id'], label=bs['label'], source=s['interpretation_reference'])
                context = dict(score=r['score_display'], band=bs['label'], answered=str(len(good)), items=str(len(use)))
                r['text'] = re.sub(r'\{\{([a-z_]+)\}\}', lambda m: context[m[1]], bs.get('text', ''))
                held = bs.get('hold', False)
                break
        if scale['bands'] and r['band'] is None:
            raise DataError('A computed score lies outside the configured bands.')
        return r, held

    def _evaluate_expression(self, node, trace, variable_trace):
        """Decimal value, None when an input is missing, or ValueError for an impossible computation."""
        if 'const' in node:
            return number(node['const'])
        if 'item' in node:
            t = trace[node['item']]
            return number(t['scored']) if t['status'] == 'valid' else None
        if 'variable' in node:
            t = variable_trace[node['variable']]
            return number(t['value']) if t['status'] == 'valid' else None
        args = [self._evaluate_expression(a, trace, variable_trace) for a in node['args']]
        if any(a is None for a in args):
            return None
        op = node['op']
        if op == 'add':
            return sum(args, Decimal(0))
        if op == 'sub':
            return args[0] - args[1]
        if op == 'mul':
            out = Decimal(1)
            for a in args:
                out *= a
            return out
        if op == 'div':
            if args[1] == 0:
                raise ValueError('Division by zero.')
            return args[0] / args[1]
        if op == 'min':
            return min(args)
        if op == 'max':
            return max(args)
        if op == 'abs':
            return abs(args[0])
        # minutes_between(start, end), wrapping past midnight. Decimal % keeps the dividend's sign, so add 1440 first.
        return ((args[1] - args[0]) % 1440 + 1440) % 1440

    def _evaluate(self, node, trace, variable_trace, by_id):
        """Three-valued condition: True, False or None (unknown)."""
        if 'all' in node or 'any' in node:
            is_all = 'all' in node
            values = [self._evaluate(n, trace, variable_trace, by_id) for n in node['all' if is_all else 'any']]
            if is_all and False in values: return False
            if not is_all and True in values: return True
            return None if None in values else is_all
        if 'not' in node:
            value = self._evaluate(node['not'], trace, variable_trace, by_id)
            return None if value is None else not value
        if 'scale' in node:
            s = by_id[node['scale']]
            value = number(s['score_exact']) if s['status'] == 'scored' else None
            rhs = number(node['value'])
        elif 'item' in node:
            t = trace[node['item']]
            value = number(t['scored']) if t['status'] == 'valid' else None
            rhs = number(node['value'])
        else:
            v = self.variables[node['variable']]; t = variable_trace[node['variable']]
            if t['status'] != 'valid':
                return None
            if v['type'] == 'categorical':
                return compare(t['value'], node['op'], key(node['value']))
            value = number(t['value'])
            rhs = clock(node['value']) if v['type'] == 'time' and isinstance(node['value'], str) else number(node['value'])
        if value is None:
            return None
        return compare(value, node['op'], rhs)

    def _fill_value(self, spec, trace, by_id):
        """Text for a section token, or None when it cannot be determined (the section is then withheld)."""
        if 'response_label' in spec:
            t = trace[spec['response_label']]
            return self.items[spec['response_label']]['labels'][t['code']] if t['status'] == 'valid' else None
        if 'score' in spec:
            s = by_id[spec['score']]
            return s['score_display'] if s['status'] == 'scored' else None
        if 'band' in spec:
            s = by_id[spec['band']]
            return s['band']['label'] if s['status'] == 'scored' and s['band'] else None
        s = by_id[spec['top_items']]
        if s['status'] != 'scored':
            return None
        answered = [i for i in s['item_ids'] if trace[i]['status'] == 'valid']
        best = max(number(trace[i]['scored']) for i in answered)
        top = [i for i in answered if number(trace[i]['scored']) == best]
        ties = spec.get('ties', 'all')
        if len(top) > 1 and ties == 'withhold':
            return None
        if ties == 'first':
            top = top[:1]
        return join_labels([self.items[i]['source']['label'] for i in top])


def score_rows(engine, rows):
    results = [engine.score(r) for r in rows]
    ids = [r['participant_id'] for r in results]
    if len(ids) != len(set(ids)):
        raise DataError('Duplicate participant ids; no reports should be written.')
    return results


def render_html(report):
    if report['status'] not in {'ready', 'partial'}:
        raise DataError(f'Report is {report["status"]}; release is withheld.')
    esc = lambda x: html.escape(str(x), quote=True)
    rows, paragraphs = [], []
    for s in report['scales']:
        score = s['score_display'] if s['status'] == 'scored' else 'Not scored'
        band = s['band']['label'] if s['band'] else 'Not assigned'
        coverage = f'{s["answered"]} / {s["expected"]}' + ('; prorated' if s['prorated'] else '')
        rows.append('<tr>' + ''.join('<td>' + esc(x) + '</td>' for x in (s['label'], score, band, coverage)) + '</tr>')
        paragraphs.append('<h2>' + esc(s['label']) + '</h2><p>' + esc(s['text'] if s['status'] == 'scored' else s['reason']) + '</p>')
    return ('<!doctype html><html lang="en"><meta charset="utf-8"><title>' + esc(report['title']) + '</title>'
            '<style>body{font:17px/1.5 Arial,sans-serif;max-width:860px;margin:40px auto;padding:0 24px;color:#171717}h1{font-size:28px}h2{font-size:21px}table{border-collapse:collapse;width:100%}th,td{border:1px solid #d9d9d9;padding:10px;text-align:left}th{background:#edf1f4}.note{font-size:14px}</style>'
            '<h1>' + esc(report['title']) + '</h1><p>Participant ' + esc(report['participant_id']) + '</p><p>' + esc(report['introduction']) + '</p>'
            '<table><thead><tr><th>Scale</th><th>Score</th><th>Band</th><th>Answered items</th></tr></thead><tbody>' + ''.join(rows) + '</tbody></table>'
            + ''.join(paragraphs) + ''.join('<h2>'+esc(s['title'])+'</h2><p>'+esc(s['text'])+'</p>' for s in report['sections'])
            + '<p class="note">Scoring reference: ' + esc(report['scoring_reference']) + '</p><p class="note">Instrument version: '
            + esc(report['instrument_version']) + '; engine ' + VERSION + '; configuration SHA-256 ' + esc(report['config_sha256']) + '</p></html>')


def render_docx(report, path):
    # Import only when Word export is requested; scoring itself has no package dependencies.
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if report['status'] not in {'ready', 'partial'}:
        raise DataError(f'Report is {report["status"]}; release is withheld.')
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        st = d.styles[name]; st.font.name = 'Arial'; st.font.color.rgb = RGBColor(0,0,0)
        for border in list(st.element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    d.styles['Normal'].font.size = Pt(11)
    d.styles['Normal'].paragraph_format.space_after = Pt(8)
    d.styles['Title'].font.size = Pt(21)
    d.styles['Heading 2'].font.size = Pt(12)
    d.add_heading(report['title'], 0)
    d.add_paragraph('Participant ' + report['participant_id'])
    d.add_paragraph(report['introduction'])
    table = d.add_table(rows=1, cols=4); table.autofit = False
    widths = [2.2, 1.0, 1.8, 1.9]
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    for c, t in zip(table.rows[0].cells, ['Scale', 'Score', 'Band', 'Answered items']): c.text=t
    for s in report['scales']:
        vals = [s['label'], s['score_display'] if s['status']=='scored' else 'Not scored',
                s['band']['label'] if s['band'] else 'Not assigned', f'{s["answered"]} / {s["expected"]}' + ('; prorated' if s['prorated'] else '')]
        for c, t in zip(table.add_row().cells, vals): c.text=str(t)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    table._tbl.tblPr.append(borders)
    for row_no,row in enumerate(table.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if row_no==0:pr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            if row_no==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF1F4');c._tc.get_or_add_tcPr().append(sh)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                v=OxmlElement('w:'+edge);v.set(qn('w:w'),'90');v.set(qn('w:type'),'dxa');margins.append(v)
            c._tc.get_or_add_tcPr().append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for run in p.runs: run.font.size=Pt(10);run.bold=row_no==0
    d.add_paragraph()
    for s in report['scales']:
        d.add_heading(s['label'],2)
        d.add_paragraph(s['text'] if s['status']=='scored' else s['reason'])
        if s['prorated']:d.add_paragraph('Proration rule: ' + s['proration_justification'])
    for s in report['sections']:
        d.add_heading(s['title'],2);d.add_paragraph(s['text'])
    d.add_paragraph('Scoring reference: ' + report['scoring_reference'])
    p=d.add_paragraph('Instrument version ' + report['instrument_version'] + '; engine ' + VERSION)
    for run in p.runs:run.font.size=Pt(9)
    p=d.add_paragraph('Configuration SHA-256: ' + report['config_sha256'])
    for run in p.runs:run.font.size=Pt(8)
    d.core_properties.author='';d.core_properties.title=report['title']
    d.save(path)


def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--config',required=True);ap.add_argument('--data',required=True)
    ap.add_argument('--out',required=True);ap.add_argument('--docx',action='store_true')
    args=ap.parse_args(argv)
    engine=ReportEngine(load_config(args.config)); data=Path(args.data).read_bytes()
    with Path(args.data).open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f)
        if reader.fieldnames is None or len(reader.fieldnames)!=len(set(reader.fieldnames)):
            raise DataError('CSV headers must be present and unique.')
        rows=list(reader)
        if any(None in r for r in rows):raise DataError('A CSV row has extra fields.')
        if any(v is None for r in rows for v in r.values()):
            raise DataError('A CSV row has fewer fields than the header; no reports were written.')
    results=score_rows(engine,rows)
    out=Path(args.out)
    if out.exists() and any(out.iterdir()):raise DataError('Output directory is not empty; choose a new directory.')
    out.mkdir(parents=True,exist_ok=True)
    counts={}
    for r in results:
        counts[r['status']]=counts.get(r['status'],0)+1
        stem='report_'+hashlib.sha256(r['participant_id'].encode()).hexdigest()[:16]
        r['output_stem']=stem
        (out/(stem+'.json')).write_text(json.dumps(r,indent=2,ensure_ascii=False),encoding='utf-8')
        if r['status'] in {'ready','partial'}:
            (out/(stem+'.html')).write_text(render_html(r),encoding='utf-8')
            if args.docx:render_docx(r,out/(stem+'.docx'))
    manifest=dict(engine_version=VERSION,config_sha256=engine.config_sha256,
                  input_sha256=hashlib.sha256(data).hexdigest(),records=len(results),status_counts=counts,
                  outputs=[dict(participant_id=r['participant_id'],status=r['status'],stem=r['output_stem']) for r in results])
    (out/'run_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    print(json.dumps(manifest,indent=2))


if __name__=='__main__':
    main()
